# import relevant libraries
import os
import re
import json
import pandas as pd
import numpy as np
import spacy
import PyPDF2
import docx
import nltk
import torch
from flask import Flask, request, jsonify, render_template
from werkzeug.utils import secure_filename
from nltk.corpus import stopwords
from nltk.tokenize import word_tokenize
from nltk.stem import WordNetLemmatizer
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from transformers import AutoTokenizer, AutoModel
from typing import List, Dict, Tuple, Any, Union, Optional

# Download necessary NLTK data
nltk.download('punkt_tab')
nltk.download('stopwords')
nltk.download('wordnet')

# Load spaCy model
nlp = spacy.load('en_core_web_sm')

# Global constants for configuration purposes
ALLOWED_EXTENSIONS = {'pdf', 'docx', 'txt', 'csv', 'json'}
UPLOADS = 'uploads'
MODEL = 'sentence-transformers/all-MiniLM-L6-v2'  # transformer model: Smaller, faster model

#Create a class for generating BERT embeding for texts
class BERTEmbedder:
    
    def __init__(self, model: str = MODEL):
        self.tokenizer = AutoTokenizer.from_pretrained(model)
        self.model = AutoModel.from_pretrained(model)
        self.device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
        self.model.to(self.device)

    # define mean pooling  function to get sentence embeddings    
    def mean_pooling(self, model_output, attention_mask):
        token_embeddings = model_output[0]
        input_mask_expanded = attention_mask.unsqueeze(-1).expand(token_embeddings.size()).float()
        return torch.sum(token_embeddings * input_mask_expanded, 1) / torch.clamp(input_mask_expanded.sum(1), min=1e-9)
    
    # define Get BERT embeddings for a list of texts
    def get_embeddings(self, texts: List[str]) -> np.ndarray:
        encoded_input = self.tokenizer(texts, padding=True, truncation=True, max_length=256, return_tensors='pt')
        encoded_input = {k: v.to(self.device) for k, v in encoded_input.items()}
        with torch.no_grad():
            model_output = self.model(**encoded_input)
        sentence_embeddings = self.mean_pooling(model_output, encoded_input['attention_mask'])
        return sentence_embeddings.cpu().numpy()


# Creating a class for Enhanced database of skills with domain classifications
class SkillsDatabase:
    
    def __init__(self, skills_file: str = None):
        # Default skills by domain if no file is provided
        self.skills_by_domain = {
            "programming_languages": [
                "python", "java", "javascript", "typescript", "c++", "c#", "ruby", "php", 
                "swift", "kotlin", "golang", "rust", "scala", "perl", "r"
            ],
            "web_development": [
                "html", "css", "sass", "react", "angular", "vue", "redux", "jquery", 
                "bootstrap", "tailwind", "webpack", "next.js", "django", "flask", "spring", 
                "laravel", "node.js", "express.js", "restful api", "graphql"
            ],
            "data_science": [
                "machine learning", "deep learning", "tensorflow", "pytorch", "scikit-learn",
                "pandas", "numpy", "data mining", "data visualization", "tableau", "power bi",
                "statistics", "r", "big data", "data modeling", "data analysis", "nlp",
                "computer vision", "neural networks", "regression", "classification"
            ],
            "database": [
                "sql", "mysql", "postgresql", "mongodb", "oracle", "sqlite", "nosql", 
                "redis", "elasticsearch", "cassandra", "mariadb", "ms sql server",
                "database design", "data warehousing", "etl", "database administration"
            ],
            "devops": [
                "aws", "azure", "gcp", "docker", "kubernetes", "jenkins", "ci/cd", "terraform",
                "ansible", "prometheus", "grafana", "git", "github", "gitlab", "bitbucket",
                "cloud computing", "microservices", "containerization", "linux", "bash"
            ],
            "soft_skills": [
                "communication", "teamwork", "leadership", "problem solving", "critical thinking",
                "time management", "adaptability", "creativity", "project management", "negotiation",
                "conflict resolution", "emotional intelligence", "decision making", "strategic thinking"
            ],
            "management": [
                "team management", "product management", "agile", "scrum", "kanban", "okr",
                "strategic planning", "budget management", "resource allocation", "stakeholder management",
                "risk management", "kpi", "business analysis", "change management", "mentoring"
            ],
            "design": [
                "ui design", "ux design", "graphic design", "photoshop", "illustrator",
                "indesign", "sketch", "figma", "adobe xd", "wireframing", "prototyping",
                "responsive design", "web design", "typography", "color theory"
            ],
            "marketing": [
                "seo", "sem", "content marketing", "social media marketing", "email marketing",
                "google analytics", "conversion optimization", "a/b testing", "market research",
                "brand management", "customer segmentation", "marketing strategy", "crm"
            ],
            "finance": [
                "financial analysis", "accounting", "budgeting", "forecasting", "financial modeling",
                "excel", "financial reporting", "taxation", "risk assessment", "investment analysis",
                "cost accounting", "balance sheet", "cash flow", "income statement"
            ]
        }
        
        # If a skills file is provided, load skills from it
        if skills_file and os.path.exists(skills_file):
            with open(skills_file, 'r') as f:
                self.skills_by_domain = json.load(f)
        
        # Create a flat list of all skills for quick lookup
        self.all_skills = []
        for domain, skills in self.skills_by_domain.items():
            for skill in skills:
                self.all_skills.append(skill)
        
        # Build regex patterns for matching skills
        self.domain_patterns = {}
        for domain, skills in self.skills_by_domain.items():
            pattern = r'\b(' + '|'.join([re.escape(skill) for skill in skills]) + r')\b'
            self.domain_patterns[domain] = re.compile(pattern, re.IGNORECASE)
        
        # Build a single pattern for all skills
        self.all_skills_pattern = re.compile(
            r'\b(' + '|'.join([re.escape(skill) for skill in self.all_skills]) + r')\b', 
            re.IGNORECASE
        )
    
    # Define a function to determine the importance of each domain for a job description
    def get_domain_importance(self, job_description: str) -> Dict[str, float]:
        domain_counts = {}
        domain_importance = {}
        
        for domain, pattern in self.domain_patterns.items():
            matches = pattern.findall(job_description.lower())
            domain_counts[domain] = len(matches)
        
        total_matches = sum(domain_counts.values()) or 1  # Avoid division by zero
        for domain, count in domain_counts.items():
            domain_importance[domain] = 0.5 + (0.5 * count / total_matches)
        
        return domain_importance
    
    def extract_skills_with_domain(self, text: str) -> Dict[str, Dict[str, float]]:
        """Extract skills from text and organize by domain with importance weights"""
        skills_by_domain = {}
        
        # Initialize empty skill lists for each domain
        for domain in self.skills_by_domain.keys():
            skills_by_domain[domain] = {}
        
        # Extract skills by domain
        for domain, pattern in self.domain_patterns.items():
            matches = pattern.findall(text.lower())
            for skill in matches:
                if skill in skills_by_domain[domain]:
                    skills_by_domain[domain][skill] += 0.1  # Increase weight for repeated mentions
                else:
                    skills_by_domain[domain][skill] = 1.0
        
        return skills_by_domain
    
    def save_skills_to_file(self, filename: str):
        """Save skills database to a JSON file"""
        with open(filename, 'w') as f:
            json.dump(self.skills_by_domain, f, indent=2)
    
    def add_skill(self, skill: str, domain: str):
        """Add a new skill to the database"""
        if domain in self.skills_by_domain:
            if skill not in self.skills_by_domain[domain]:
                self.skills_by_domain[domain].append(skill)
                self.all_skills.append(skill)
                # Rebuild the patterns
                pattern = r'\b(' + '|'.join([re.escape(s) for s in self.skills_by_domain[domain]]) + r')\b'
                self.domain_patterns[domain] = re.compile(pattern, re.IGNORECASE)
                self.all_skills_pattern = re.compile(
                    r'\b(' + '|'.join([re.escape(s) for s in self.all_skills]) + r')\b', 
                    re.IGNORECASE
                )
                return True
        return False


class ExperienceLevelEstimator:
    """Class for estimating candidate experience level from resume"""
    
    def __init__(self):
        # Keywords associated with different experience levels
        self.junior_keywords = [
            "junior", "entry level", "entry-level", "intern", "internship", "assistant",
            "fresh graduate", "recent graduate", "0-2 years", "less than 2 years"
        ]
        self.mid_keywords = [
            "mid level", "mid-level", "intermediate", "associate", "2-5 years", 
            "3-5 years", "experienced"
        ]
        self.senior_keywords = [
            "senior", "lead", "principal", "architect", "manager", "head of", "director",
            "5+ years", "5-10 years", "6+ years", "7+ years", "expert", "advanced"
        ]
        
        # Education level indicators
        self.education_levels = {
            "high school": 1,
            "associate degree": 2,
            "bachelor": 3,
            "master": 4,
            "phd": 5,
            "doctorate": 5,
            "mba": 4
        }
    
    def estimate_level(self, resume_text: str) -> Dict[str, Any]:
        """Estimate the experience level based on resume text"""
        result = {
            "level": "unknown",
            "confidence": 0.0,
            "years_experience": 0,
            "education_level": 0
        }
        
        # Convert to lowercase for case-insensitive matching
        text = resume_text.lower()
        
        # Detect experience level from keywords
        junior_score = sum(1 for kw in self.junior_keywords if kw in text)
        mid_score = sum(1 for kw in self.mid_keywords if kw in text)
        senior_score = sum(1 for kw in self.senior_keywords if kw in text)
        
        # Determine highest score
        scores = [
            ("junior", junior_score),
            ("mid", mid_score),
            ("senior", senior_score)
        ]
        
        # Sort by score in descending order
        scores.sort(key=lambda x: x[1], reverse=True)
        
        # If highest score is positive, set as estimated level
        if scores[0][1] > 0:
            result["level"] = scores[0][0]
            total_score = junior_score + mid_score + senior_score
            if total_score > 0:
                result["confidence"] = scores[0][1] / total_score
        
        # Extract years of experience using regex
        years_patterns = [
            r'(\d+)\+?\s*years?\s*(?:of)?\s*experience',
            r'experience\s*(?:of)?\s*(\d+)\+?\s*years?',
            r'worked\s*(?:for)?\s*(\d+)\+?\s*years?'
        ]
        
        for pattern in years_patterns:
            matches = re.findall(pattern, text)
            if matches:
                years = max([int(y) for y in matches])
                result["years_experience"] = years
                
                # Adjust level based on years if confidence is low
                if result["confidence"] < 0.6:
                    if years < 3:
                        result["level"] = "junior"
                    elif years < 6:
                        result["level"] = "mid"
                    else:
                        result["level"] = "senior"
                break
        
        # Detect highest education level
        for edu_level, value in self.education_levels.items():
            if edu_level in text:
                result["education_level"] = max(result["education_level"], value)
        
        return result


class JobRequirementAnalyzer:
    """Class for analyzing job descriptions and extracting structured requirements"""
    
    def __init__(self, skills_db: SkillsDatabase):
        self.skills_db = skills_db
        
        # Patterns for extracting different types of requirements
        self.experience_pattern = re.compile(r'(\d+)[\+]?\s+years?\s+(?:of\s+)?experience', re.IGNORECASE)
        self.education_pattern = re.compile(
            r'(bachelor|master|phd|doctorate|mba|bs|ms|ba|degree|diploma)', 
            re.IGNORECASE
        )
        self.required_pattern = re.compile(
            r'(required|requirements?|must have|essential|necessary)', 
            re.IGNORECASE
        )
        self.preferred_pattern = re.compile(
            r'(preferred|nice to have|desirable|plus|advantage|beneficial)', 
            re.IGNORECASE
        )
    
    def analyze(self, job_description: str) -> Dict[str, Any]:
        """Extract structured requirements from job description"""
        # Split job description into sentences
        doc = nlp(job_description)
        sentences = [sent.text for sent in doc.sents]
        
        result = {
            "required_skills": {},
            "preferred_skills": {},
            "min_years_experience": 0,
            "education_requirements": [],
            "domain_importance": {}
        }
        
        # Extract domain importance
        result["domain_importance"] = self.skills_db.get_domain_importance(job_description)
        
        # Process each sentence
        for sentence in sentences:
            sentence_lower = sentence.lower()
            
            # Check if sentence contains required or preferred keywords
            is_required = bool(self.required_pattern.search(sentence_lower))
            is_preferred = bool(self.preferred_pattern.search(sentence_lower))
            
            # Skip sentences that don't indicate requirements
            if not (is_required or is_preferred):
                continue
            
            # Extract skills from the sentence
            skills_in_domain = self.skills_db.extract_skills_with_domain(sentence)
            
            # Add skills to appropriate category
            for domain, skills in skills_in_domain.items():
                if not skills:
                    continue
                    
                if is_required and not is_preferred:
                    # Add to required skills
                    if domain not in result["required_skills"]:
                        result["required_skills"][domain] = {}
                    result["required_skills"][domain].update(skills)
                else:
                    # Add to preferred skills
                    if domain not in result["preferred_skills"]:
                        result["preferred_skills"][domain] = {}
                    result["preferred_skills"][domain].update(skills)
            
            # Extract years of experience
            experience_matches = self.experience_pattern.findall(sentence)
            if experience_matches:
                years = max([int(x) for x in experience_matches])
                result["min_years_experience"] = max(result["min_years_experience"], years)
            
            # Extract education requirements
            education_matches = self.education_pattern.findall(sentence)
            if education_matches:
                for edu in education_matches:
                    if edu.lower() not in result["education_requirements"]:
                        result["education_requirements"].append(edu.lower())
        
        # If no explicit required skills were found, use all extracted skills
        if not result["required_skills"]:
            all_skills = self.skills_db.extract_skills_with_domain(job_description)
            result["required_skills"] = all_skills
        
        return result


class DocumentProcessor:
    """Enhanced class for processing different document formats"""
    
    def extract_text_from_file(self, file_path: str) -> str:
        """Extract text from different file formats"""
        file_extension = os.path.splitext(file_path)[1].lower()
        
        if file_extension == '.pdf':
            return self.extract_text_from_pdf(file_path)
        elif file_extension == '.docx':
            return self.extract_text_from_docx(file_path)
        elif file_extension == '.txt':
            return self.extract_text_from_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {file_extension}")
    
    def extract_text_from_pdf(self, pdf_path: str) -> str:
        """Extract text from PDF files"""
        text = ""
        try:
            with open(pdf_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page_num in range(len(pdf_reader.pages)):
                    page = pdf_reader.pages[page_num]
                    text += page.extract_text()
        except Exception as e:
            print(f"Error extracting text from PDF: {e}")
            return ""
        return text
    
    def extract_text_from_docx(self, docx_path: str) -> str:
        """Extract text from DOCX files"""
        try:
            doc = docx.Document(docx_path)
            full_text = []
            
            # Extract text from paragraphs
            for para in doc.paragraphs:
                full_text.append(para.text)
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        full_text.append(cell.text)
            
            return '\n'.join(full_text)
        except Exception as e:
            print(f"Error extracting text from DOCX: {e}")
            return ""
    
    def extract_text_from_txt(self, txt_path: str) -> str:
        """Extract text from TXT files"""
        try:
            with open(txt_path, 'r', encoding='utf-8', errors='ignore') as file:
                return file.read()
        except Exception as e:
            print(f"Error extracting text from TXT: {e}")
            return ""
    
    # def load_resume_dataset(self, dataset_path: str, format_type: str = "csv") -> Tuple[List[str], List[str]]:
    #     """Load multiple resumes from a dataset file
        
    #     Args:
    #         dataset_path: Path to the dataset file
    #         format_type: Format of the dataset (csv or json)
            
    #     Returns:
    #         Tuple containing list of resume texts and list of resume IDs
    #     """
    #     resume_texts = []
    #     resume_ids = []
        
    #     try:
    #         if format_type.lower() == "csv":
    #             import pandas as pd
    #             df = pd.read_csv(dataset_path)
    #             # Identify resume text column (look for column with 'resume' in name)
    #             resume_columns = [col for col in df.columns if 'resume' in col.lower() or 'text' in col.lower()]
    #             if resume_columns:
    #                 resume_column = resume_columns[0]
    #             else:
    #                 # If no obvious column, use the column with the longest text on average
    #                 text_lengths = {col: df[col].astype(str).str.len().mean() for col in df.columns}
    #                 resume_column = max(text_lengths, key=text_lengths.get)
    #                 print(f"Using column '{resume_column}' as resume text column")
                
    #             resume_texts = df[resume_column].astype(str).tolist()
    #             # Use index or another column as ID
    #             if 'id' in df.columns:
    #                 resume_ids = df['id'].astype(str).tolist()
    #             else:
    #                 resume_ids = [f"resume_{i}" for i in range(len(resume_texts))]
            
    #         elif format_type.lower() == "json":
    #             import json
    #             with open(dataset_path, 'r', encoding='utf-8') as f:
    #                 data = json.load(f)
                    
    #             # Handle different possible JSON structures
    #             if isinstance(data, list):
    #                 # List of dictionaries
    #                 if all(isinstance(item, dict) for item in data):
    #                     # Look for keys containing 'resume' or 'text'
    #                     for item in data:
    #                         text_key = next((k for k in item.keys() if 'resume' in k.lower() or 'text' in k.lower()), None)
    #                         if text_key:
    #                             resume_texts.append(str(item[text_key]))
    #                         else:
    #                             # Use the value from the longest text field
    #                             text_fields = {k: len(str(v)) for k, v in item.items() if isinstance(v, (str, int, float))}
    #                             if text_fields:
    #                                 longest_field = max(text_fields, key=text_fields.get)
    #                                 resume_texts.append(str(item[longest_field]))
    #                             else:
    #                                 resume_texts.append("")  # Empty if no suitable field
                                    
    #                         # Try to find an ID field
    #                         id_key = next((k for k in item.keys() if 'id' in k.lower()), None)
    #                         if id_key:
    #                             resume_ids.append(str(item[id_key]))
    #                         else:
    #                             resume_ids.append(f"resume_{len(resume_ids)}")
    #             else:
    #                 raise ValueError("JSON format not recognized. Expected a list of resume objects.")
            
    #         else:
    #             raise ValueError(f"Unsupported format type: {format_type}. Use 'csv' or 'json'.")
                
    #         print(f"Successfully loaded {len(resume_texts)} resumes from dataset")
    #         return resume_texts, resume_ids
            
    #     except Exception as e:
    #         print(f"Error loading resume dataset: {str(e)}")
    #         return [], []

    def load_resume_dataset(self, dataset_path: str, format_type: str = "csv") -> Tuple[List[str], List[str]]:
        """Load multiple resumes from a dataset file"""
        resume_texts = []
        resume_ids = []
        
        try:
            if format_type.lower() == "csv":
                import pandas as pd
                print(f"Loading CSV file from {dataset_path}")
                
                # Try different encodings if one fails
                try:
                    df = pd.read_csv(dataset_path, encoding='utf-8')
                except UnicodeDecodeError:
                    df = pd.read_csv(dataset_path, encoding='latin1')
                
                print(f"CSV loaded successfully with {len(df)} rows")
                print(f"Columns found: {list(df.columns)}")
                
                # Identify the resume text column - look for specific column names
                resume_column = None
                potential_columns = ['resume_text', 'resume', 'text', 'content', 'description']
                
                # Try exact matches first
                for col in potential_columns:
                    if col in df.columns:
                        resume_column = col
                        break
                
                # If no exact match, try partial matches
                if resume_column is None:
                    for col in df.columns:
                        if any(potential in col.lower() for potential in potential_columns):
                            resume_column = col
                            break
                
                # If still no match, use the column with the longest text
                if resume_column is None:
                    text_lengths = {col: df[col].astype(str).str.len().mean() for col in df.columns}
                    resume_column = max(text_lengths, key=text_lengths.get)
                
                print(f"Using column '{resume_column}' for resume text")
                
                # Process each row individually
                for idx, row in df.iterrows():
                    # Get resume text
                    text = str(row[resume_column])
                    if len(text.strip()) > 0:
                        resume_texts.append(text)
                        
                        # Determine ID (try common ID fields or use index)
                        if 'id' in df.columns:
                            resume_ids.append(str(row['id']))
                        elif 'ID' in df.columns:
                            resume_ids.append(str(row['ID']))
                        else:
                            resume_ids.append(f"Resume_{idx+1}")
                
                print(f"Successfully extracted {len(resume_texts)} resumes from dataset")
                
            elif format_type.lower() == "json":
                # [existing JSON handling code]
                pass
                
            else:
                raise ValueError(f"Unsupported format type: {format_type}. Use 'csv' or 'json'.")
            
            if not resume_texts:
                print("WARNING: No resume texts were extracted from the file!")
            else:
                print(f"First resume (sample): {resume_texts[0][:100]}...")
                
            return resume_texts, resume_ids
            
        except Exception as e:
            print(f"Error loading resume dataset: {str(e)}")
            import traceback
            traceback.print_exc()
            return [], []


class EnhancedResumeRanker:
    """Enhanced class for ranking resumes against job descriptions"""
    
    def __init__(self, use_transformer: bool = True):
        self.stop_words = set(stopwords.words('english'))
        self.lemmatizer = WordNetLemmatizer()
        self.document_processor = DocumentProcessor()
        self.skills_db = SkillsDatabase()
        self.experience_estimator = ExperienceLevelEstimator()
        self.requirement_analyzer = JobRequirementAnalyzer(self.skills_db)
        self.use_transformer = use_transformer
        
        # Initialize BERT embedder for semantic matching
        if use_transformer:
            self.bert_embedder = BERTEmbedder()
        
        # Initialize TF-IDF vectorizer as fallback
        self.vectorizer = TfidfVectorizer(
            min_df=1,
            max_df=0.85,
            ngram_range=(1, 2),
            sublinear_tf=True
        )
        
        # Default weights for scoring components
        self.default_weights = {
            "semantic_similarity": 0.35,
            "required_skills_match": 0.30,
            "preferred_skills_match": 0.15,
            "experience_match": 0.15,
            "education_match": 0.05
        }
        
        # Job-type specific weights
        self.job_type_weights = {
            "technical": {
                "semantic_similarity": 0.30,
                "required_skills_match": 0.40,
                "preferred_skills_match": 0.15,
                "experience_match": 0.10,
                "education_match": 0.05
            },
            "management": {
                "semantic_similarity": 0.25,
                "required_skills_match": 0.20,
                "preferred_skills_match": 0.15,
                "experience_match": 0.30,
                "education_match": 0.10
            },
            "entry_level": {
                "semantic_similarity": 0.40,
                "required_skills_match": 0.25,
                "preferred_skills_match": 0.10,
                "experience_match": 0.05,
                "education_match": 0.20
            }
        }
    
    def preprocess_text(self, text: str) -> str:
        """Clean and preprocess text data"""
        # Convert to lowercase
        text = text.lower()
        
        # Remove special characters and numbers
        text = re.sub(r'[^\w\s]', ' ', text)
        text = re.sub(r'\d+', ' ', text)
        
        # Tokenize
        tokens = word_tokenize(text)
        
        # Remove stopwords and lemmatize
        processed_tokens = [
            self.lemmatizer.lemmatize(token) 
            for token in tokens 
            if token not in self.stop_words and len(token) > 2
        ]
        
        return ' '.join(processed_tokens)
    
    def calculate_semantic_similarity(self, job_description: str, 
                                     resumes: List[str]) -> np.ndarray:
        """Calculate semantic similarity using BERT or TF-IDF"""
        if self.use_transformer:
            # Use BERT embeddings
            all_texts = [job_description] + resumes
            embeddings = self.bert_embedder.get_embeddings(all_texts)
            
            # Calculate cosine similarity
            job_embedding = embeddings[0].reshape(1, -1)
            resume_embeddings = embeddings[1:]
            
            similarity_scores = cosine_similarity(job_embedding, resume_embeddings)
            return similarity_scores[0]
        else:
            # Fallback to TF-IDF
            all_documents = [job_description] + resumes
            tfidf_matrix = self.vectorizer.fit_transform(all_documents)
            
            job_vector = tfidf_matrix[0]
            resume_vectors = tfidf_matrix[1:]
            
            similarity_scores = cosine_similarity(job_vector, resume_vectors)
            return similarity_scores[0]
    
    def calculate_skills_match(self, job_requirements: Dict[str, Any], 
                              resume_text: str, 
                              required: bool = True) -> Dict[str, float]:
        """Calculate how well a resume matches job required or preferred skills"""
        # Extract skills from resume by domain
        resume_skills = self.skills_db.extract_skills_with_domain(resume_text)
        
        # Get job skills (either required or preferred)
        job_skills = job_requirements["required_skills"] if required else job_requirements["preferred_skills"]
        
        # Get domain importance weights
        domain_importance = job_requirements["domain_importance"]
        
        # Initialize results
        results = {
            "overall_match_score": 0.0,
            "domain_scores": {},
            "matched_skills": {},
            "missing_skills": {}
        }
        
        # Calculate match by domain
        total_weight = 0
        weighted_sum = 0
        
        for domain, skills in job_skills.items():
            domain_weight = domain_importance.get(domain, 0.5)
            total_weight += domain_weight
            
            # Initialize domain results
            results["domain_scores"][domain] = 0.0
            results["matched_skills"][domain] = []
            results["missing_skills"][domain] = []
            
            if not skills:
                continue
                
            # Count matches
            matches = 0
            for skill, importance in skills.items():
                if domain in resume_skills and skill in resume_skills[domain]:
                    matches += 1
                    results["matched_skills"][domain].append(skill)
                else:
                    results["missing_skills"][domain].append(skill)
            
            # Calculate domain score (percentage of skills matched)
            domain_score = matches / len(skills) if skills else 0
            results["domain_scores"][domain] = domain_score
            
            # Add to weighted sum
            weighted_sum += domain_score * domain_weight
        
        # Calculate overall match score
        if total_weight > 0:
            results["overall_match_score"] = weighted_sum / total_weight
        
        return results
    
    def evaluate_experience_match(self, job_requirements: Dict[str, Any], 
                                 resume_text: str) -> Dict[str, Any]:
        """Evaluate how well a candidate's experience matches job requirements"""
        # Estimate candidate experience level
        experience_info = self.experience_estimator.estimate_level(resume_text)
        
        # Get minimum required years
        min_years = job_requirements.get("min_years_experience", 0)
        
        result = {
            "candidate_years": experience_info["years_experience"],
            "required_years": min_years,
            "candidate_level": experience_info["level"],
            "match_score": 0.0
        }
        
        # Calculate experience match score
        if min_years > 0 and experience_info["years_experience"] > 0:
            if experience_info["years_experience"] >= min_years:
                # Candidate meets or exceeds requirements
                result["match_score"] = min(1.0, experience_info["years_experience"] / (min_years * 1.5))
            else:
                # Candidate has some but insufficient experience
                result["match_score"] = experience_info["years_experience"] / min_years
        elif min_years == 0:
            # No specific experience requirement
            result["match_score"] = 1.0
        elif experience_info["level"] != "unknown":
            # No specific years detected, use level as proxy
            level_scores = {"junior": 0.3, "mid": 0.7, "senior": 1.0}
            result["match_score"] = level_scores.get(experience_info["level"], 0.0)
        
        return result
    
    def evaluate_education_match(self, job_requirements: Dict[str, Any], 
                               resume_text: str) -> Dict[str, Any]:
        """Evaluate how well a candidate's education matches job requirements"""
        # Get required education from job
        required_education = job_requirements.get("education_requirements", [])
        
        # Map education terms to levels
        education_levels = {
            "high school": 1,
            "associate": 2,
            "bachelor": 3, "bs": 3, "ba": 3, "undergraduate": 3,
            "master": 4, "ms": 4, "ma": 4, "graduate": 4,
            "phd": 5, "doctorate": 5,
            "mba": 4
        }
        
        # Find highest required education level
        required_level = 0
        for edu in required_education:
            for term, level in education_levels.items():
                if term in edu.lower():
                    required_level = max(required_level, level)
        
        # Estimate candidate's education level
        candidate_level = 0
        for term, level in education_levels.items():
            if term in resume_text.lower():
                candidate_level = max(candidate_level, level)
        
        result = {
            "candidate_education_level": candidate_level,
            "required_education_level": required_level,
            "match_score": 0.0
        }
        
        # Calculate education match score
        if required_level == 0:
            # No specific education requirement
            result["match_score"] = 1.0
        elif candidate_level >= required_level:
            # Candidate meets or exceeds requirements
            result["match_score"] = 1.0
        else:
            # Candidate has lower education level
            result["match_score"] = candidate_level / required_level
        
        return result
    
    def determine_job_type(self, job_description: str, job_requirements: Dict[str, Any]) -> str:
        """Determine the job type for weight customization"""
        # Look for management keywords
        management_keywords = ["manager", "director", "lead", "supervisor", "chief", "head of"]
        is_management = any(keyword in job_description.lower() for keyword in management_keywords)
        
        # Check for entry level indicators
        entry_keywords = ["entry level", "junior", "internship", "intern", "trainee", "graduate"]
        min_years = job_requirements.get("min_years_experience", 0)
        is_entry_level = (min_years <= 1) or any(keyword in job_description.lower() for keyword in entry_keywords)
        
        # Determine type
        if is_management:
            return "management"
        elif is_entry_level:
            return "entry_level"
        else:
            return "technical"
    
    def get_weights_for_job(self, job_description: str, job_requirements: Dict[str, Any]) -> Dict[str, float]:
        """Get the appropriate scoring weights based on job type"""
        job_type = self.determine_job_type(job_description, job_requirements)
        
        if job_type in self.job_type_weights:
            return self.job_type_weights[job_type]
        
        return self.default_weights
    
    def rank_resumes(self, job_description: str, resume_files: List[str], 
                    detailed: bool = True) -> List[Dict[str, Any]]:
        """Rank resumes based on their relevance to the job description"""
        # Extract job requirements
        job_requirements = self.requirement_analyzer.analyze(job_description)
        
        # Get appropriate weights for this job
        weights = self.get_weights_for_job(job_description, job_requirements)
        
        # Extract resumes from files
        resume_texts = []
        for resume_file in resume_files:
            resume_text = self.document_processor.extract_text_from_file(resume_file)
            resume_texts.append(resume_text)
        
        # Preprocess job description and resumes
        processed_job = self.preprocess_text(job_description)
        processed_resumes = [self.preprocess_text(resume) for resume in resume_texts]
        
        # Calculate semantic similarity
        similarity_scores = self.calculate_semantic_similarity(processed_job, processed_resumes)
        
        # Prepare ranking results
        ranking_results = []
        
        for i, (resume_file, resume_text, similarity) in enumerate(
            zip(resume_files, resume_texts, similarity_scores)
        ):
            # Base result with file info
            result = {
                "file_name": os.path.basename(resume_file),
                "semantic_similarity": float(similarity),
                "rank": i + 1  # Will be updated later
            }
            
            if detailed:
                # Calculate required skills match
                required_match = self.calculate_skills_match(
                    job_requirements, resume_text, required=True
                )
                result["required_skills_match"] = required_match
                
                # Calculate preferred skills match
                preferred_match = self.calculate_skills_match(
                    job_requirements, resume_text, required=False
                )
                result["preferred_skills_match"] = preferred_match
                
                # Evaluate experience match
                experience_match = self.evaluate_experience_match(job_requirements, resume_text)
                result["experience_match"] = experience_match
                
                # Evaluate education match
                education_match = self.evaluate_education_match(job_requirements, resume_text)
                result["education_match"] = education_match
                
                # Calculate final score (weighted combination)
                result["component_scores"] = {
                    "semantic_similarity": similarity * weights["semantic_similarity"],
                    "required_skills_match": required_match["overall_match_score"] * weights["required_skills_match"],
                    "preferred_skills_match": preferred_match["overall_match_score"] * weights["preferred_skills_match"],
                    "experience_match": experience_match["match_score"] * weights["experience_match"],
                    "education_match": education_match["match_score"] * weights["education_match"]
                }
                
                # Sum all component scores
                result["final_score"] = sum(result["component_scores"].values())
                
                # Store summary for display
                result["summary"] = {
                    "matched_required_skills": sum(len(skills) for skills in required_match["matched_skills"].values()),
                    "total_required_skills": sum(len(skills) for skills in job_requirements["required_skills"].values()),
                    "matched_preferred_skills": sum(len(skills) for skills in preferred_match["matched_skills"].values()),
                    "total_preferred_skills": sum(len(skills) for skills in job_requirements["preferred_skills"].values()),
                    "experience_years": experience_match["candidate_years"],
                    "required_years": experience_match["required_years"],
                    "education_level": education_match["candidate_education_level"],
                    "weights_used": weights
                }
            else:
                result["final_score"] = float(similarity)
            
            ranking_results.append(result)
        
        # Sort by final score
        ranking_results = sorted(ranking_results, key=lambda x: x["final_score"], reverse=True)
        
        # Update ranks
        for i, result in enumerate(ranking_results):
            result["rank"] = i + 1
        
        return ranking_results
    
    def rank_dataset_resumes(self, job_description: str, dataset_path: str, format_type: str = "csv", detailed: bool = True) -> List[Dict[str, Any]]:
        
        # Extract job requirements
        job_requirements = self.requirement_analyzer.analyze(job_description)
        
        # Get appropriate weights for this job
        weights = self.get_weights_for_job(job_description, job_requirements)
        
        # Load resumes from dataset
        resume_texts, resume_ids = self.document_processor.load_resume_dataset(dataset_path, format_type)
        
        if not resume_texts:
            return []
        
        # Preprocess job description and resumes
        processed_job = self.preprocess_text(job_description)
        processed_resumes = [self.preprocess_text(resume) for resume in resume_texts]
        
        # Calculate semantic similarity
        similarity_scores = self.calculate_semantic_similarity(processed_job, processed_resumes)
        
        # Prepare ranking results
        ranking_results = []
        
        for i, (resume_id, resume_text, similarity) in enumerate(
            zip(resume_ids, resume_texts, similarity_scores)
        ):
            # Base result with resume info
            result = {
                "resume_id": resume_id,
                "semantic_similarity": float(similarity),
                "rank": i + 1  # Will be updated later
            }
            
            # Continue with the same logic as in rank_resumes method...
            # Just replacing file_name with resume_id

            if detailed:
            # Calculate required skills match
                required_match = self.calculate_skills_match(
                    job_requirements, resume_text, required=True
                )
                result["required_skills_match"] = required_match
                
                # Calculate preferred skills match
                preferred_match = self.calculate_skills_match(
                    job_requirements, resume_text, required=False
                )
                result["preferred_skills_match"] = preferred_match
                
                # Evaluate experience match
                experience_match = self.evaluate_experience_match(job_requirements, resume_text)
                result["experience_match"] = experience_match
                
                # Evaluate education match
                education_match = self.evaluate_education_match(job_requirements, resume_text)
                result["education_match"] = education_match
                
                # Calculate final score (weighted combination)
                result["component_scores"] = {
                    "semantic_similarity": similarity * weights["semantic_similarity"],
                    "required_skills_match": required_match["overall_match_score"] * weights["required_skills_match"],
                    "preferred_skills_match": preferred_match["overall_match_score"] * weights["preferred_skills_match"],
                    "experience_match": experience_match["match_score"] * weights["experience_match"],
                    "education_match": education_match["match_score"] * weights["education_match"]
                }
                
                # Sum all component scores
                result["final_score"] = sum(result["component_scores"].values())
                
                # Store summary for display
                result["summary"] = {
                    "matched_required_skills": sum(len(skills) for skills in required_match["matched_skills"].values()),
                    "total_required_skills": sum(len(skills) for skills in job_requirements["required_skills"].values()),
                    "matched_preferred_skills": sum(len(skills) for skills in preferred_match["matched_skills"].values()),
                    "total_preferred_skills": sum(len(skills) for skills in job_requirements["preferred_skills"].values()),
                    "experience_years": experience_match["candidate_years"],
                    "required_years": experience_match["required_years"],
                    "education_level": education_match["candidate_education_level"],
                    "weights_used": weights
                }
        else:
            # If not detailed, just use similarity as final score
            result["final_score"] = float(similarity)
        
            
            ranking_results.append(result)
        
        # Sort by final score
        ranking_results = sorted(ranking_results, key=lambda x: x["final_score"], reverse=True)
        
        # Update ranks
        for i, result in enumerate(ranking_results):
            result["rank"] = i + 1
        
        return ranking_results
    
    def get_top_ranked_resumes(self, job_description: str, dataset_path: str, format_type: str = "csv", max_results: int = 20) -> List[Dict[str, Any]]:
        # Get all ranked resumes
        all_results = self.rank_dataset_resumes(job_description, dataset_path, format_type)
        
        # Results should already be sorted by rank/final_score, but let's make sure
        sorted_results = sorted(all_results, key=lambda x: x["final_score"], reverse=True)
        
        # If we have more than max_results, limit to just the top ones
        if len(sorted_results) > max_results:
            print(f"Found {len(sorted_results)} resumes, limiting to top {max_results}")
            return sorted_results[:max_results]
        
        # Otherwise return all results
        return sorted_results    


# Create a Flask web application
app = Flask(__name__, template_folder='templates')
app.config['UPLOADS'] = UPLOADS
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16 MB max upload size

# Ensure upload folder exists
os.makedirs(UPLOADS, exist_ok=True)

def allowed_file(filename):
    """Check if file has an allowed extension"""
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@app.route('/')
def index():
    """Render the main page"""
    return render_template('index.html')

# @app.route('/api/rank', methods=['POST'])
# def rank_api():
#     """API endpoint for resume ranking"""
#     try:
#         print("Received request to /api/rank")
#         print(f"Form keys: {list(request.form.keys())}")
#         print(f"Files keys: {list(request.files.keys())}")
#         print("Content type:", request.content_type)
        
#         # Check if job description was provided
#         if 'job_description' not in request.form:
#             return jsonify({"error": "No job description provided"}), 400
        
#         job_description = request.form['job_description']
#         print(f"Job description length: {len(job_description)}")
        
#         # Create resume ranker
#         ranker = EnhancedResumeRanker()

#         # Make sure upload folder exists
#         os.makedirs(app.config['UPLOADS'], exist_ok=True)

#         # Check for any dataset file (try multiple possible field names)
#         dataset_file = None
#         dataset_field_names = ['datasetFile', 'dataset', 'dataset_file']
#         for field_name in dataset_field_names:
#             if field_name in request.files and request.files[field_name].filename:
#                 dataset_file = request.files[field_name]
#                 print(f"Found dataset file in field '{field_name}': {dataset_file.filename}")
#                 break
        
#         # Check for any resume files (try multiple possible field names)
#         resume_files = []
#         resume_field_names = ['resumeFiles', 'resumes', 'resume_files']
#         for field_name in resume_field_names:
#             if field_name in request.files:
#                 files = request.files.getlist(field_name)
#                 resume_files.extend([f for f in files if f.filename])
#                 if resume_files:
#                     print(f"Found {len(resume_files)} resume files in field '{field_name}'")
#                     break
        
#         # Determine if we're using dataset or individual files
#         if dataset_file:
#             # Dataset mode
#             print("Processing dataset file upload")
#             dataset_file = request.files['datasetFile']
#             filename = dataset_file.filename
#             print(f"Dataset filename: {filename}")
            
#             # Check file extension
#             if not ('.' in filename and filename.rsplit('.', 1)[1].lower() in ['csv', 'json']):
#                 return jsonify({"error": "Invalid file format. Please upload a CSV or JSON file"}), 400
            
#             # Determine format from file extension if not specified in form
#             if 'datasetFormat' in request.form:
#                 dataset_format = request.form['datasetFormat']
#             else:
#                 dataset_format = filename.rsplit('.', 1)[1].lower()
            
#             print(f"Dataset format: {dataset_format}")
            
#             # Save dataset file temporarily
#             dataset_filename = secure_filename(dataset_file.filename)
#             dataset_path = os.path.join(app.config['UPLOADS'], dataset_filename)
#             dataset_file.save(dataset_path)
#             print(f"Dataset saved to: {dataset_path}")
            
#             try:
#                 # Ensure upload folder exists
#                 os.makedirs(app.config['UPLOADS'], exist_ok=True)
                
#                 # Rank resumes from dataset
#                 print("Ranking resumes from dataset...")
#                 results = ranker.rank_dataset_resumes(job_description, dataset_path, dataset_format)
#                 print(f"Ranking complete. Found {len(results)} results")
                
#                 # Analyze job requirements
#                 job_requirements = ranker.requirement_analyzer.analyze(job_description)
                
#                 # Clean up temporary file
#                 if os.path.exists(dataset_path):
#                     os.remove(dataset_path)
#                     print(f"Deleted temporary file: {dataset_path}")
                
#                 return jsonify({"results": results, "job_analysis": job_requirements})
            
#             except Exception as e:
#                 print(f"Error processing dataset: {str(e)}")
#                 # Clean up temporary file
#                 if os.path.exists(dataset_path):
#                     os.remove(dataset_path)
                
#                 return jsonify({"error": f"Error processing dataset: {str(e)}"}), 500
        
#         elif resume_files:
#             # Individual files mode
#             print("Processing individual resume files")
#             files = request.files.getlist('resumeFiles')
#             print(f"Number of resume files: {len(files)}")
            
#             # Save uploaded files
#             resume_paths = []
#             for file in files:
#                 if file and file.filename:
#                     if allowed_file(file.filename):
#                         filename = secure_filename(file.filename)
#                         file_path = os.path.join(app.config['UPLOADS'], filename)
#                         # Ensure directory exists
#                         os.makedirs(os.path.dirname(file_path), exist_ok=True)
#                         file.save(file_path)
#                         resume_paths.append(file_path)
#                         print(f"Saved resume file: {file_path}")
#                     else:
#                         print(f"Skipping file with invalid format: {file.filename}")
            
#             if not resume_paths:
#                 return jsonify({"error": "No valid resume files uploaded. Please upload PDF, DOCX, or TXT files."}), 400
            
#             # Rank resumes
#             try:
#                 print("Ranking individual resumes...")
#                 results = ranker.rank_resumes(job_description, resume_paths)
#                 print(f"Ranking complete. Found {len(results)} results")
                
#                 # Analyze job requirements
#                 job_requirements = ranker.requirement_analyzer.analyze(job_description)
                
#                 # Clean up temporary files
#                 for path in resume_paths:
#                     if os.path.exists(path):
#                         os.remove(path)
#                         print(f"Deleted temporary file: {path}")
                
#                 return jsonify({"results": results, "job_analysis": job_requirements})
            
#             except Exception as e:
#                 print(f"Error ranking resumes: {str(e)}")
#                 # Clean up temporary files
#                 for path in resume_paths:
#                     if os.path.exists(path):
#                         os.remove(path)
                
#                 return jsonify({"error": f"Error ranking resumes: {str(e)}"}), 500
        
#         else:
#             print("No resume files or dataset uploaded")
#             return jsonify({"error": "No resume files or dataset uploaded. Please upload either individual resume files or a dataset file."}), 400
    
#     except Exception as e:
#         print(f"Unexpected error in rank_api: {str(e)}")
#         import traceback
#         traceback.print_exc()
#         return jsonify({"error": f"An unexpected error occurred: {str(e)}"}), 500
    
@app.route('/api/rank', methods=['POST'])
def rank_api():
    """API endpoint for resume ranking that processes files without relying on specific field names"""
    try:
        print("Received request to /api/rank")
        print(f"Form keys: {list(request.form.keys())}")
        print(f"Files keys: {list(request.files.keys())}")
        
        # Check if job description was provided
        if 'job_description' not in request.form:
            return jsonify({"error": "No job description provided"}), 400
        
        job_description = request.form['job_description']
        print(f"Job description length: {len(job_description)}")
        
        # Create resume ranker
        ranker = EnhancedResumeRanker()
        
        # Make sure upload folder exists
        os.makedirs(app.config['UPLOADS'], exist_ok=True)
        
        # Collect all files from all fields
        all_files = []
        for field_name in request.files:
            file_list = request.files.getlist(field_name)
            all_files.extend([f for f in file_list if f.filename])
        
        print(f"Found {len(all_files)} total files in the request")
        
        # Categorize files by extension
        dataset_file = None
        dataset_format = None
        resume_files = []
        
        for file in all_files:
            filename = file.filename.lower()
            extension = filename.rsplit('.', 1)[1] if '.' in filename else ''
            
            if extension in ['csv', 'json']:
                # This looks like a dataset file
                if not dataset_file:  # Take the first dataset file we find
                    dataset_file = file
                    dataset_format = extension
                    print(f"Found dataset file: {filename} (format: {extension})")
            elif extension in ['pdf', 'docx', 'txt']:
                # This looks like a resume file
                resume_files.append(file)
                print(f"Found resume file: {filename}")
            else:
                print(f"Ignoring file with unsupported extension: {filename}")
        
        # Check if dataset format is specified in form (override file extension if provided)
        if 'datasetFormat' in request.form and dataset_file:
            dataset_format = request.form['datasetFormat']
            print(f"Using dataset format from form: {dataset_format}")
        
        # Process based on what was found
        if dataset_file:
            # Dataset mode
            print("Processing dataset file")
            
            # Save dataset file temporarily
            dataset_filename = secure_filename(dataset_file.filename)
            dataset_path = os.path.join(app.config['UPLOADS'], dataset_filename)
            dataset_file.save(dataset_path)
            print(f"Dataset saved to: {dataset_path}")
            
            try:
                # Rank resumes from dataset
                print(f"Ranking resumes from dataset ({dataset_format})...")
                results = ranker.rank_dataset_resumes(job_description, dataset_path, dataset_format)
                print(f"Ranking complete. Found {len(results)} results")
                
                # Analyze job requirements
                job_requirements = ranker.requirement_analyzer.analyze(job_description)
                
                # Clean up temporary file
                if os.path.exists(dataset_path):
                    os.remove(dataset_path)
                    print(f"Deleted temporary file: {dataset_path}")
                
                return jsonify({"results": results, "job_analysis": job_requirements})
            
            except Exception as e:
                print(f"Error processing dataset: {str(e)}")
                # Clean up temporary file
                if os.path.exists(dataset_path):
                    os.remove(dataset_path)
                
                return jsonify({"error": f"Error processing dataset: {str(e)}"}), 500
        
        elif resume_files:
            # Individual files mode
            print(f"Processing {len(resume_files)} individual resume files")
            
            # Save uploaded files
            resume_paths = []
            for file in resume_files:
                if file and file.filename:
                    filename = secure_filename(file.filename)
                    file_path = os.path.join(app.config['UPLOADS'], filename)
                    # Ensure directory exists
                    os.makedirs(os.path.dirname(file_path), exist_ok=True)
                    file.save(file_path)
                    resume_paths.append(file_path)
                    print(f"Saved resume file: {file_path}")
            
            if not resume_paths:
                return jsonify({"error": "No valid resume files could be saved."}), 400
            
            # Rank resumes
            try:
                print("Ranking individual resumes...")
                results = ranker.rank_resumes(job_description, resume_paths)
                print(f"Ranking complete. Found {len(results)} results")
                
                # Analyze job requirements
                job_requirements = ranker.requirement_analyzer.analyze(job_description)
                
                # Clean up temporary files
                for path in resume_paths:
                    if os.path.exists(path):
                        os.remove(path)
                        print(f"Deleted temporary file: {path}")
                
                return jsonify({"results": results, "job_analysis": job_requirements})
            
            except Exception as e:
                print(f"Error ranking resumes: {str(e)}")
                # Clean up temporary files
                for path in resume_paths:
                    if os.path.exists(path):
                        os.remove(path)
                
                return jsonify({"error": f"Error ranking resumes: {str(e)}"}), 500
        
        else:
            print("No resume files or dataset uploaded")
            return jsonify({"error": "No resume files or dataset uploaded. Please upload either resume files (PDF, DOCX, TXT) or a dataset file (CSV, JSON)."}), 400
            
    except Exception as e:
        print(f"Unexpected error in rank_api: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({"error": f"An unexpected error occurred: {str(e)}"}), 500
    

@app.route('/api/analyze-job', methods=['POST'])
def analyze_job_api():
    """API endpoint for job description analysis"""
    # Check if job description was provided
    if 'job_description' not in request.form:
        return jsonify({"error": "No job description provided"}), 400
    
    job_description = request.form['job_description']
    
    # Create requirement analyzer
    skills_db = SkillsDatabase()
    analyzer = JobRequirementAnalyzer(skills_db)
    
    # Analyze job description
    try:
        job_requirements = analyzer.analyze(job_description)
        return jsonify({"job_analysis": job_requirements})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

def main():
    """Main function for command-line usage"""
    import argparse
    
    # Create the parser
    parser = argparse.ArgumentParser(description='Resume Ranking System')
    parser.add_argument('--job', type=str, help='Path to job description file')
    parser.add_argument('--resumes', type=str, nargs='+', help='Paths to resume files')
    parser.add_argument('--dataset', type=str, help='Path to dataset file with multiple resumes')
    parser.add_argument('--format', type=str, default='csv', choices=['csv', 'json'], help='Format of dataset file (csv or json)')
    parser.add_argument('--web', action='store_true', help='Start web interface')
    parser.add_argument('--port', type=int, default=5000, help='Port for web interface')
    
    # Parse arguments
    args = parser.parse_args()
    
    if args.web:
        # Start web interface
        print(f"Starting web interface at http://localhost:{args.port}")
        print("Press Ctrl+C to stop the server")
        app.run(host='0.0.0.0', port=args.port, debug=False)
    elif args.job and args.resumes:
        # Command-line mode
        # Read job description
        try:
            with open(args.job, 'r', encoding='utf-8') as f:
                job_description = f.read()
        except UnicodeDecodeError:
            # Fall back to a more permissive encoding
            with open(args.job, 'r', encoding='latin-1') as f:  # or 'cp1252' for Windows
                job_description = f.read()
        
        # Create resume ranker
        ranker = EnhancedResumeRanker()
        
        # Analyze job description
        job_requirements = ranker.requirement_analyzer.analyze(job_description)
        print("\nJob Analysis:")
        print(f"Required Skills: {job_requirements['required_skills']}")
        print(f"Preferred Skills: {job_requirements['preferred_skills']}")
        print(f"Minimum Experience: {job_requirements['min_years_experience']} years")
        print(f"Education Requirements: {job_requirements['education_requirements']}")
        print(f"Domain Importance: {job_requirements['domain_importance']}")
        
        # Rank resumes
        print("\nRanking Resumes...")
        results = ranker.rank_resumes(job_description, args.resumes)
        
        # Display results
        print("\nRanking Results:")
        for result in results:
            print(f"\nRank {result['rank']}: {result['file_name']} (Score: {result['final_score']:.2f})")
            print(f"  Semantic Similarity: {result['component_scores']['semantic_similarity']:.2f}")
            print(f"  Required Skills Match: {result['component_scores']['required_skills_match']:.2f}")
            print(f"  Preferred Skills Match: {result['component_scores']['preferred_skills_match']:.2f}")
            print(f"  Experience Match: {result['component_scores']['experience_match']:.2f}")
            print(f"  Education Match: {result['component_scores']['education_match']:.2f}")
            
            print(f"  Matched {result['summary']['matched_required_skills']}/{result['summary']['total_required_skills']} required skills")
            print(f"  Matched {result['summary']['matched_preferred_skills']}/{result['summary']['total_preferred_skills']} preferred skills")
            print(f"  Experience: {result['summary']['experience_years']} years (Required: {result['summary']['required_years']})")
    
    elif args.job and args.dataset:
        # Read job description
        with open(args.job, 'r', encoding='utf-8') as f:
            job_description = f.read()
        
        # Create resume ranker
        ranker = EnhancedResumeRanker()
        
        # Rank resumes from dataset
        print(f"\nRanking resumes from dataset {args.dataset}...")
        results = ranker.rank_dataset_resumes(job_description, args.dataset, args.format)
        
        # Display results
        print("\nRanking Results:")
        for result in results:
            print(f"\nRank {result['rank']}: Resume ID {result['resume_id']} (Score: {result['final_score']:.2f})")
            # Display other details similar to the individual file ranking output
    
    else:
        # No valid arguments provided
        parser.print_help()

# HTML template for the web interface
@app.route('/templates')
def get_template():
    return """
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Resume Ranking System</title>
    <link href="https://cdn.jsdelivr.net/npm/tailwindcss@2.2.19/dist/tailwind.min.css" rel="stylesheet">
</head>
<body class="bg-gray-100 min-h-screen">
    <div class="container mx-auto px-4 py-8">
        <header class="mb-8">
            <h1 class="text-3xl font-bold text-gray-800">AI Resume Ranking System</h1>
            <p class="text-gray-600">Upload a job description and resumes to find the best matches</p>
        </header>
        
        <div class="bg-white rounded-lg shadow-md p-6 mb-8">
            <h2 class="text-xl font-semibold mb-4">Upload Files</h2>
            <form id="rankingForm" class="space-y-4">
                <div>
                    <label class="block text-gray-700 mb-2" for="jobDescription">Job Description</label>
                    <textarea id="jobDescription" name="jobDescription" rows="6" 
                        class="w-full px-3 py-2 border border-gray-300 rounded-md focus:outline-none focus:ring-2 focus:ring-blue-500"
                        placeholder="Paste the job description here..."></textarea>
                </div>
                
                <div>
                    <label class="block text-gray-700 mb-2" for="resumeFiles">Resume Files</label>
                    <input type="file" id="resumeFiles" name="resumeFiles" multiple
                        class="w-full px-3 py-2 border border-gray-300 rounded-md focus:outline-none focus:ring-2 focus:ring-blue-500"
                        accept=".pdf,.docx,.txt">
                    <p class="text-sm text-gray-500 mt-1">Upload PDF, DOCX, or TXT files</p>
                </div>
                
                <div class="flex justify-between">
                    <button type="button" id="analyzeJobBtn" 
                        class="px-4 py-2 bg-gray-600 text-white rounded-md hover:bg-gray-700 focus:outline-none focus:ring-2 focus:ring-gray-500">
                        Analyze Job Only
                    </button>
                    <button type="submit" 
                        class="px-4 py-2 bg-blue-600 text-white rounded-md hover:bg-blue-700 focus:outline-none focus:ring-2 focus:ring-blue-500">
                        Rank Resumes
                    </button>
                </div>
            </form>
        </div>
        
        <div id="jobAnalysisSection" class="bg-white rounded-lg shadow-md p-6 mb-8 hidden">
            <h2 class="text-xl font-semibold mb-4">Job Analysis</h2>
            <div id="jobAnalysisContent" class="space-y-4"></div>
        </div>
        
        <div id="resultsSection" class="bg-white rounded-lg shadow-md p-6 hidden">
            <h2 class="text-xl font-semibold mb-4">Ranking Results</h2>
            <div id="resultsContent"></div>
        </div>
        
        <div id="loadingSection" class="text-center py-8 hidden">
            <div class="inline-block animate-spin rounded-full h-8 w-8 border-b-2 border-gray-900"></div>
            <p class="mt-2 text-gray-600">Processing...</p>
        </div>
    </div>
    
    <script>
        document.addEventListener('DOMContentLoaded', function() {
            const rankingForm = document.getElementById('rankingForm');
            const analyzeJobBtn = document.getElementById('analyzeJobBtn');
            const jobAnalysisSection = document.getElementById('jobAnalysisSection');
            const jobAnalysisContent = document.getElementById('jobAnalysisContent');
            const resultsSection = document.getElementById('resultsSection');
            const resultsContent = document.getElementById('resultsContent');
            const loadingSection = document.getElementById('loadingSection');
            
            // Handle form submission for full ranking
            rankingForm.addEventListener('submit', async function(e) {
                e.preventDefault();
                
                const jobDescription = document.getElementById('jobDescription').value;
                const resumeFiles = document.getElementById('resumeFiles').files;
                
                if (!jobDescription) {
                    alert('Please enter a job description');
                    return;
                }
                
                if (resumeFiles.length === 0) {
                    alert('Please upload at least one resume file');
                    return;
                }
                
                // Show loading indicator
                loadingSection.classList.remove('hidden');
                jobAnalysisSection.classList.add('hidden');
                resultsSection.classList.add('hidden');
                
                // Create form data
                const formData = new FormData();
                formData.append('job_description', jobDescription);
                
                for (let i = 0; i < resumeFiles.length; i++) {
                    formData.append('resumes', resumeFiles[i]);
                }
                
                try {
                    // Send request to API
                    const response = await fetch('/api/rank', {
                        method: 'POST',
                        body: formData
                    });
                    
                    const data = await response.json();
                    
                    if (response.ok) {
                        // Display job analysis
                        displayJobAnalysis(data.job_analysis);
                        
                        // Display ranking results
                        displayResults(data.results);
                    } else {
                        alert(`Error: ${data.error}`);
                    }
                } catch (error) {
                    alert(`Error: ${error.message}`);
                } finally {
                    // Hide loading indicator
                    loadingSection.classList.add('hidden');
                }
            });
            
            // Handle job analysis only
            analyzeJobBtn.addEventListener('click', async function() {
                const jobDescription = document.getElementById('jobDescription').value;
                
                if (!jobDescription) {
                    alert('Please enter a job description');
                    return;
                }
                
                // Show loading indicator
                loadingSection.classList.remove('hidden');
                jobAnalysisSection.classList.add('hidden');
                resultsSection.classList.add('hidden');
                
                // Create form data
                const formData = new FormData();
                formData.append('job_description', jobDescription);
                
                try {
                    // Send request to API
                    const response = await fetch('/api/analyze-job', {
                        method: 'POST',
                        body: formData
                    });
                    
                    const data = await response.json();
                    
                    if (response.ok) {
                        // Display job analysis
                        displayJobAnalysis(data.job_analysis);
                    } else {
                        alert(`Error: ${data.error}`);
                    }
                } catch (error) {
                    alert(`Error: ${error.message}`);
                } finally {
                    // Hide loading indicator
                    loadingSection.classList.add('hidden');
                }
            });
            
            // Function to display job analysis
            function displayJobAnalysis(jobAnalysis) {
                jobAnalysisContent.innerHTML = '';
                
                // Create required skills section
                const requiredSkillsSection = document.createElement('div');
                requiredSkillsSection.innerHTML = `
                    <h3 class="font-medium text-gray-800">Required Skills</h3>
                    <div class="mt-2 grid grid-cols-2 gap-2">
                        ${Object.entries(jobAnalysis.required_skills).map(([domain, skills]) => `
                            <div class="p-2 border rounded">
                                <p class="font-medium">${domain.replace('_', ' ')}</p>
                                <p>${Object.keys(skills).join(', ')}</p>
                            </div>
                        `).join('')}
                    </div>
                `;
                
                // Create preferred skills section
                const preferredSkillsSection = document.createElement('div');
                preferredSkillsSection.innerHTML = `
                    <h3 class="font-medium text-gray-800 mt-4">Preferred Skills</h3>
                    <div class="mt-2 grid grid-cols-2 gap-2">
                        ${Object.entries(jobAnalysis.preferred_skills).map(([domain, skills]) => `
                            <div class="p-2 border rounded">
                                <p class="font-medium">${domain.replace('_', ' ')}</p>
                                <p>${Object.keys(skills).join(', ')}</p>
                            </div>
                        `).join('')}
                    </div>
                `;
                
                // Create requirements section
                const requirementsSection = document.createElement('div');
                requirementsSection.innerHTML = `
                    <h3 class="font-medium text-gray-800 mt-4">Requirements</h3>
                    <div class="mt-2 space-y-2">
                        <p><strong>Experience:</strong> ${jobAnalysis.min_years_experience} years</p>
                        <p><strong>Education:</strong> ${jobAnalysis.education_requirements.join(', ') || 'Not specified'}</p>
                    </div>
                `;
                
                // Create domain importance section
                const domainSection = document.createElement('div');
                domainSection.innerHTML = `
                    <h3 class="font-medium text-gray-800 mt-4">Domain Importance</h3>
                    <div class="mt-2 grid grid-cols-2 gap-2">
                        ${Object.entries(jobAnalysis.domain_importance).map(([domain, importance]) => `
                            <div class="p-2 border rounded">
                                <p class="font-medium">${domain.replace('_', ' ')}</p>
                                <div class="w-full bg-gray-200 rounded-full h-2.5">
                                    <div class="bg-blue-600 h-2.5 rounded-full" style="width: ${importance * 100}%"></div>
                                </div>
                            </div>
                        `).join('')}
                    </div>
                `;
                
                // Add all sections to content
                jobAnalysisContent.appendChild(requiredSkillsSection);
                jobAnalysisContent.appendChild(preferredSkillsSection);
                jobAnalysisContent.appendChild(requirementsSection);
                jobAnalysisContent.appendChild(domainSection);
                
                // Show job analysis section
                jobAnalysisSection.classList.remove('hidden');
            }
            
            // Function to display results
            function displayResults(results) {
                resultsContent.innerHTML = '';
                
                if (results.length === 0) {
                    resultsContent.innerHTML = '<p>No results found</p>';
                    resultsSection.classList.remove('hidden');
                    return;
                }
                
                // Create table
                const table = document.createElement('table');
                table.className = 'min-w-full divide-y divide-gray-200';
                
                // Create table header
                const thead = document.createElement('thead');
                thead.className = 'bg-gray-50';
                thead.innerHTML = `
                    <tr>
                        <th scope="col" class="px-6 py-3 text-left text-xs font-medium text-gray-500 uppercase tracking-wider">Rank</th>
                        <th scope="col" class="px-6 py-3 text-left text-xs font-medium text-gray-500 uppercase tracking-wider">Resume</th>
                        <th scope="col" class="px-6 py-3 text-left text-xs font-medium text-gray-500 uppercase tracking-wider">Overall Score</th>
                        <th scope="col" class="px-6 py-3 text-left text-xs font-medium text-gray-500 uppercase tracking-wider">Skills Match</th>
                        <th scope="col" class="px-6 py-3 text-left text-xs font-medium text-gray-500 uppercase tracking-wider">Experience</th>
                        <th scope="col" class="px-6 py-3 text-left text-xs font-medium text-gray-500 uppercase tracking-wider">Details</th>
                    </tr>
                `;
                
                // Create table body
                const tbody = document.createElement('tbody');
                tbody.className = 'bg-white divide-y divide-gray-200';
                
                results.forEach((result, index) => {
                    const tr = document.createElement('tr');
                    if (index % 2 === 0) {
                        tr.className = 'bg-white';
                    } else {
                        tr.className = 'bg-gray-50';
                    }
                    
                    const requiredSkillsMatch = result.summary.matched_required_skills / 
                                             (result.summary.total_required_skills || 1);
                    
                    tr.innerHTML = `
                        <td class="px-6 py-4 whitespace-nowrap text-sm font-medium text-gray-900">${result.rank}</td>
                        <td class="px-6 py-4 whitespace-nowrap text-sm text-gray-500">${result.file_name}</td>
                        <td class="px-6 py-4 whitespace-nowrap text-sm text-gray-500">
                            <div class="flex items-center">
                                <span class="mr-2 font-medium">${(result.final_score * 100).toFixed(0)}%</span>
                                <div class="w-full bg-gray-200 rounded-full h-2.5">
                                    <div class="bg-blue-600 h-2.5 rounded-full" style="width: ${result.final_score * 100}%"></div>
                                </div>
                            </div>
                        </td>
                        <td class="px-6 py-4 whitespace-nowrap text-sm text-gray-500">
                            ${result.summary.matched_required_skills}/${result.summary.total_required_skills}
                            <div class="w-full bg-gray-200 rounded-full h-2.5">
                                <div class="bg-green-600 h-2.5 rounded-full" style="width: ${requiredSkillsMatch * 100}%"></div>
                            </div>
                        </td>
                        <td class="px-6 py-4 whitespace-nowrap text-sm text-gray-500">
                            ${result.summary.experience_years} yrs
                            <span class="text-xs">(Required: ${result.summary.required_years})</span>
                        </td>
                        <td class="px-6 py-4 whitespace-nowrap text-right text-sm font-medium">
                            <button class="text-blue-600 hover:text-blue-900 show-details-btn" data-index="${index}">
                                Show Details
                            </button>
                        </td>
                    `;
                    
                    tbody.appendChild(tr);
                    
                    // Create details row (hidden by default)
                    const detailsRow = document.createElement('tr');
                    detailsRow.className = 'details-row hidden';
                    detailsRow.setAttribute('data-index', index);
                    
                    const detailsCell = document.createElement('td');
                    detailsCell.colSpan = 6;
                    detailsCell.className = 'px-6 py-4';
                    
                    // Component scores
                    const componentScores = document.createElement('div');
                    componentScores.className = 'grid grid-cols-5 gap-4 mb-4';
                    
                    for (const [component, score] of Object.entries(result.component_scores)) {
                        const scoreDiv = document.createElement('div');
                        scoreDiv.innerHTML = `
                            <p class="text-sm font-medium">${formatComponentName(component)}</p>
                            <div class="flex items-center mt-1">
                                <span class="mr-2 text-sm">${(score * 10).toFixed(1)}</span>
                                <div class="w-full bg-gray-200 rounded-full h-2">
                                    <div class="bg-indigo-600 h-2 rounded-full" style="width: ${score * 10 * 10}%"></div>
                                </div>
                            </div>
                        `;
                        componentScores.appendChild(scoreDiv);
                    }
                    
                    // Matched skills
                    const matchedSkills = document.createElement('div');
                    matchedSkills.className = 'mt-4';
                    matchedSkills.innerHTML = `
                        <h4 class="text-sm font-medium mb-2">Matched Required Skills</h4>
                        <div class="flex flex-wrap gap-2">
                            ${Object.entries(result.required_skills_match.matched_skills).flatMap(([domain, skills]) => 
                                skills.map(skill => `
                                    <span class="px-2 py-1 bg-green-100 text-green-800 text-xs rounded-full">
                                        ${skill}
                                    </span>
                                `)
                            ).join('')}
                        </div>
                        
                        <h4 class="text-sm font-medium mt-4 mb-2">Missing Required Skills</h4>
                        <div class="flex flex-wrap gap-2">
                            ${Object.entries(result.required_skills_match.missing_skills).flatMap(([domain, skills]) => 
                                skills.map(skill => `
                                    <span class="px-2 py-1 bg-red-100 text-red-800 text-xs rounded-full">
                                        ${skill}
                                    </span>
                                `)
                            ).join('')}
                        </div>
                    `;
                    
                    detailsCell.appendChild(componentScores);
                    detailsCell.appendChild(matchedSkills);
                    detailsRow.appendChild(detailsCell);
                    tbody.appendChild(detailsRow);
                });
                
                table.appendChild(thead);
                table.appendChild(tbody);
                resultsContent.appendChild(table);
                
                // Add event listeners for show details buttons
                document.querySelectorAll('.show-details-btn').forEach(button => {
                    button.addEventListener('click', function() {
                        const index = this.getAttribute('data-index');
                        const detailsRow = document.querySelector(`.details-row[data-index="${index}"]`);
                        
                        if (detailsRow.classList.contains('hidden')) {
                            // Hide all detail rows
                            document.querySelectorAll('.details-row').forEach(row => {
                                row.classList.add('hidden');
                            });
                            
                            // Show this detail row
                            detailsRow.classList.remove('hidden');
                            this.textContent = 'Hide Details';
                        } else {
                            detailsRow.classList.add('hidden');
                            this.textContent = 'Show Details';
                        }
                    });
                });
                
                // Show results section
                resultsSection.classList.remove('hidden');
            }
            
            // Helper function to format component names
            function formatComponentName(name) {
                return name
                    .replace(/_/g, ' ')
                    .split(' ')
                    .map(word => word.charAt(0).toUpperCase() + word.slice(1))
                    .join(' ');
            }
        });
    </script>
</body>
</html>
    """

if __name__ == "__main__":
    main()